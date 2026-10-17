from pathlib import Path

import openpyxl
from PySide6.QtWidgets import QFileDialog, QMessageBox

from pcc.core import export_xlsx
from pcc.gui.main_window import MainWindow
from pcc.gui.state import AppState


def make(qtbot, with_plate=True):
    state = AppState(legacy_paths=[])
    if with_plate:
        state.add_plate()
        state.update_plate(state.job.plates[0].id, hours=1.0)
    w = MainWindow(state)
    qtbot.addWidget(w)
    return state, w


def no_desktop(monkeypatch):
    opened = []
    monkeypatch.setattr("pcc.gui.export_flow.QDesktopServices.openUrl",
                        lambda url: opened.append(url.toLocalFile()))
    return opened


def test_export_new_then_append_last(qtbot, cfg_dir, tmp_path, monkeypatch):
    opened = no_desktop(monkeypatch)
    state, w = make(qtbot)
    flow = w.summary.export_flow
    flow.menu().aboutToShow.emit()
    assert not flow.last_action.isEnabled()
    p = tmp_path / "r.xlsx"
    assert flow.export(p, append=False) is True
    assert state.settings.last_export_path == str(p)
    assert w.toast.label.text() == "✅ Saved to r.xlsx"
    flow.menu().aboutToShow.emit()
    assert flow.last_action.isEnabled()
    assert "r.xlsx" in flow.last_action.text()
    flow.last_action.trigger()
    wb = openpyxl.load_workbook(p)
    assert wb["Jobs"].max_row == 3
    from PySide6.QtWidgets import QPushButton
    btns = {b.text(): b for b in w.toast.findChildren(QPushButton)}
    btns["Open file"].click()
    assert opened == [str(p)]
    btns["Open folder"].click()
    assert opened == [str(p), str(p.parent)]


def test_permission_error_retry(qtbot, cfg_dir, tmp_path, monkeypatch):
    state, w = make(qtbot)
    real, calls = export_xlsx.export_quote, []

    def flaky(*a, **k):
        calls.append(1)
        if len(calls) == 1:
            raise PermissionError("locked")
        return real(*a, **k)

    monkeypatch.setattr(export_xlsx, "export_quote", flaky)
    shown = []

    def fake_exec(self):
        shown.append((self.windowTitle(), self.text(), self.icon()))
        return QMessageBox.StandardButton.Retry

    monkeypatch.setattr(QMessageBox, "exec", fake_exec)
    assert w.summary.export_flow.export(tmp_path / "x.xlsx", False) is True
    assert len(calls) == 2
    assert shown == [("File is open", "Close the file in Excel, then retry",
                      QMessageBox.Icon.Warning)]


def test_permission_error_cancel(qtbot, cfg_dir, tmp_path, monkeypatch):
    state, w = make(qtbot)
    monkeypatch.setattr(export_xlsx, "export_quote",
                        lambda *a, **k: (_ for _ in ()).throw(PermissionError()))
    monkeypatch.setattr(QMessageBox, "exec", lambda self: QMessageBox.StandardButton.Cancel)
    assert w.summary.export_flow.export(tmp_path / "x.xlsx", False) is False
    assert state.settings.last_export_path == ""


def test_invalid_workbook_critical(qtbot, cfg_dir, tmp_path, monkeypatch):
    state, w = make(qtbot)
    bad = tmp_path / "bad.xlsx"
    bad.write_text("not a zip")
    msgs = []
    monkeypatch.setattr(QMessageBox, "critical", lambda p, t, m: msgs.append(m))
    assert w.summary.export_flow.export(bad, True) is False
    assert msgs and msgs[0].startswith("Could not save the report: ")


def test_export_disabled_without_quote(qtbot, cfg_dir, monkeypatch):
    state, w = make(qtbot, with_plate=False)
    assert not w.summary.export_button.isEnabled()
    called = []
    monkeypatch.setattr(w.summary.export_flow, "_ask_save_path", lambda: called.append(1) or "")
    w.summary.export_flow.quick_export()
    assert called == []
    state.add_plate()
    state.update_plate(state.job.plates[0].id, hours=1.0)
    assert w.summary.export_button.isEnabled()


def test_dialog_actions(qtbot, cfg_dir, tmp_path, monkeypatch):
    no_desktop(monkeypatch)
    state, w = make(qtbot)
    flow = w.summary.export_flow
    p = tmp_path / "n.xlsx"
    monkeypatch.setattr(flow, "_ask_save_path", lambda: str(p))
    flow.new_action.trigger()
    assert p.exists()
    monkeypatch.setattr(QFileDialog, "getOpenFileName",
                        lambda parent, t, d, flt: (str(p), flt))
    flow.append_action.trigger()
    assert openpyxl.load_workbook(p)["Jobs"].max_row == 3


def test_ctrl_e_picks_new_or_last(qtbot, cfg_dir, tmp_path, monkeypatch):
    no_desktop(monkeypatch)
    state, w = make(qtbot)
    flow = w.summary.export_flow
    calls = []
    monkeypatch.setattr(flow, "new_report", lambda: calls.append("new"))
    monkeypatch.setattr(flow, "append_last", lambda: calls.append("last"))
    flow.quick_export()
    state.update_settings(last_export_path=str(tmp_path / "gone.xlsx"))
    flow.quick_export()
    f = tmp_path / "there.xlsx"
    f.write_bytes(b"")
    state.update_settings(last_export_path=str(f))
    flow.quick_export()
    assert calls == ["new", "new", "last"]


def test_save_dialog_appends_xlsx_suffix(qtbot, cfg_dir, monkeypatch):
    state, w = make(qtbot)
    seen = {}

    def fake_exec(self):
        seen.update(suffix=self.defaultSuffix(), mode=self.acceptMode(),
                    filters=self.nameFilters(), title=self.windowTitle())
        return QFileDialog.DialogCode.Rejected

    monkeypatch.setattr(QFileDialog, "exec", fake_exec)
    assert w.summary.export_flow.new_report() is False
    assert seen == {"suffix": "xlsx", "mode": QFileDialog.AcceptMode.AcceptSave,
                    "filters": ["Excel (*.xlsx)"], "title": "Save report"}


def _overwrite_setup(qtbot, tmp_path, monkeypatch, answer):
    no_desktop(monkeypatch)
    state, w = make(qtbot)
    flow = w.summary.export_flow
    target = tmp_path / "report.xlsx"
    target.write_bytes(b"precious")
    # The dialog (e.g. a portal on Linux) returns the name without the suffix.
    monkeypatch.setattr(flow, "_ask_save_path", lambda: str(tmp_path / "report"), raising=False)
    monkeypatch.setattr(QFileDialog, "getSaveFileName",
                        lambda *a, **k: (str(tmp_path / "report"), "Excel (*.xlsx)"))
    asked = []

    def question(parent, title, text, *a, **k):
        asked.append(text)
        return answer

    monkeypatch.setattr(QMessageBox, "question", question)
    return flow, target, asked


def test_new_report_confirms_overwrite_of_suffixed_path(qtbot, cfg_dir, tmp_path, monkeypatch):
    flow, target, asked = _overwrite_setup(
        qtbot, tmp_path, monkeypatch, QMessageBox.StandardButton.No)
    assert flow.new_report() is False
    assert target.read_bytes() == b"precious"
    assert asked == ["report.xlsx already exists. Replace it?"]


def test_new_report_overwrite_confirmed(qtbot, cfg_dir, tmp_path, monkeypatch):
    flow, target, asked = _overwrite_setup(
        qtbot, tmp_path, monkeypatch, QMessageBox.StandardButton.Yes)
    assert flow.new_report() is True
    assert openpyxl.load_workbook(target)["Jobs"].max_row == 2
    assert len(asked) == 1


def test_new_report_no_extra_prompt_when_dialog_confirmed(qtbot, cfg_dir, tmp_path, monkeypatch):
    no_desktop(monkeypatch)
    state, w = make(qtbot)
    flow = w.summary.export_flow
    target = tmp_path / "r.xlsx"
    target.write_bytes(b"old")
    monkeypatch.setattr(flow, "_ask_save_path", lambda: str(target))
    monkeypatch.setattr(QMessageBox, "question",
                        lambda *a, **k: (_ for _ in ()).throw(AssertionError("asked")))
    assert flow.new_report() is True
