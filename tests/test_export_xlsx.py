from datetime import datetime

import openpyxl
import pytest

from pcc.core.export_xlsx import JOBS_HEADERS, PLATES_HEADERS, export_quote


def test_new_report(tmp_path, job, quote, catalog):
    out = export_quote(tmp_path / "r", job, quote, catalog, append=False)
    assert out.suffix == ".xlsx"
    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == ["Jobs", "Plates"]
    assert [c.value for c in wb["Jobs"][1]] == JOBS_HEADERS
    assert [c.value for c in wb["Plates"][1]] == PLATES_HEADERS
    assert wb["Jobs"]["Q2"].value == 280
    assert wb["Jobs"].freeze_panes == "A2"
    assert wb["Jobs"]["A1"].font.bold
    assert wb["Jobs"]["B2"].value == datetime(2026, 10, 17, 12, 0)
    assert wb["Jobs"]["B2"].number_format == "yyyy-mm-dd hh:mm"
    assert wb["Jobs"]["H2"].value == 7.0
    assert wb["Jobs"]["I2"].value == 84.0
    p = wb["Plates"]
    assert p["B2"].value == 1
    assert p["H2"].value == "PLA 42.0 g"
    assert p["I2"].value == pytest.approx(quote.plates[0].per_copy.filament * 2)
    assert p["N2"].value == pytest.approx(quote.plates[0].price_total)


def test_append_adds_rows(tmp_path, job, quote, catalog):
    path = tmp_path / "r.xlsx"
    export_quote(path, job, quote, catalog, append=False)
    export_quote(path, job, quote, catalog, append=True)
    wb = openpyxl.load_workbook(path)
    assert wb["Jobs"].max_row == 3
    assert wb["Plates"].max_row == 3
    export_quote(path, job, quote, catalog, append=False)
    assert openpyxl.load_workbook(path)["Jobs"].max_row == 2


def test_no_private_columns(tmp_path, job, quote, catalog):
    for h in JOBS_HEADERS + PLATES_HEADERS:
        assert not any(w in h.casefold() for w in ("margin", "fee", "per kg"))


def test_legacy_v01_sheet_kept(tmp_path, job, quote, catalog):
    path = tmp_path / "old.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Date", "Name", "Duration", "Cost"])
    ws.append(["2026-01-01", "x", 2, 10])
    wb.save(path)
    export_quote(path, job, quote, catalog, append=True)
    wb = openpyxl.load_workbook(path)
    assert wb.sheetnames == ["Legacy v0.1", "Jobs", "Plates"]
    assert [c.value for c in wb["Legacy v0.1"][2]] == ["2026-01-01", "x", 2, 10]


def test_append_to_plain_workbook_drops_empty_sheet(tmp_path, job, quote, catalog):
    path = tmp_path / "e.xlsx"
    openpyxl.Workbook().save(path)
    export_quote(path, job, quote, catalog, append=True)
    assert openpyxl.load_workbook(path).sheetnames == ["Jobs", "Plates"]


def test_unicode_customer(tmp_path, job, quote, catalog):
    job.customer = "Niño 🎉 Café"
    out = export_quote(tmp_path / "u.xlsx", job, quote, catalog, append=False)
    ws = openpyxl.load_workbook(out)["Jobs"]
    assert ws["C2"].value == "Niño 🎉 Café"
    assert ws["Q2"].number_format == '"₱"#,##0.00'


def test_locked_file_raises(tmp_path, job, quote, catalog, monkeypatch):
    def boom(self, *a, **k):
        raise PermissionError("locked")

    monkeypatch.setattr(openpyxl.Workbook, "save", boom)
    with pytest.raises(PermissionError):
        export_quote(tmp_path / "l.xlsx", job, quote, catalog, append=False)


def test_text_cells_never_formulas(tmp_path, job, quote, catalog):
    job.customer = '=HYPERLINK("http://evil","x")'
    job.plates[0].name = "=1+1"
    job.plates[0].source_file = "=cmd|' /C calc'!A0"
    catalog.printers[0].name = "@SUM(1)"
    from pcc.core.pricing import compute_quote
    from pcc.core.models import Settings
    q = compute_quote(job, Settings(), catalog)
    out = export_quote(tmp_path / "f.xlsx", job, q, catalog, append=False)
    wb = openpyxl.load_workbook(out)
    jobs, plates = wb["Jobs"], wb["Plates"]
    assert jobs["C2"].data_type == "s"
    assert jobs["C2"].value == '=HYPERLINK("http://evil","x")'
    assert jobs["D2"].data_type == "s" and jobs["D2"].value == "@SUM(1)"
    assert plates["C2"].data_type == "s" and plates["C2"].value == "=1+1"
    assert plates["D2"].data_type == "s" and plates["D2"].value == "=cmd|' /C calc'!A0"


def test_failed_save_keeps_original_intact(tmp_path, job, quote, catalog, monkeypatch):
    path = tmp_path / "keep.xlsx"
    export_quote(path, job, quote, catalog, append=False)
    before = path.read_bytes()

    def crash(self, filename):
        with open(filename, "wb") as fh:
            fh.write(b"partial")
        raise OSError("disk full")

    monkeypatch.setattr(openpyxl.Workbook, "save", crash)
    with pytest.raises(OSError):
        export_quote(path, job, quote, catalog, append=True)
    assert path.read_bytes() == before
    assert sorted(p.name for p in tmp_path.iterdir()) == ["keep.xlsx"]


def test_locked_target_propagates_and_cleans_temp(tmp_path, job, quote, catalog, monkeypatch):
    import os
    path = tmp_path / "lock.xlsx"
    export_quote(path, job, quote, catalog, append=False)
    before = path.read_bytes()

    def locked(src, dst):
        raise PermissionError("locked")

    monkeypatch.setattr(os, "replace", locked)
    with pytest.raises(PermissionError):
        export_quote(path, job, quote, catalog, append=True)
    assert path.read_bytes() == before
    assert sorted(p.name for p in tmp_path.iterdir()) == ["lock.xlsx"]


def test_foreign_workbook_backed_up_once(tmp_path, job, quote, catalog):
    path = tmp_path / "mine.xlsx"
    wb = openpyxl.Workbook()
    wb.active.append(["my", "data"])
    wb.save(path)
    original = path.read_bytes()
    export_quote(path, job, quote, catalog, append=True)
    bak = tmp_path / "mine.bak.xlsx"
    assert bak.read_bytes() == original
    # Appending again: it is a pcc workbook now and the backup is never overwritten.
    bak.write_bytes(b"sentinel")
    export_quote(path, job, quote, catalog, append=True)
    assert bak.read_bytes() == b"sentinel"


def test_foreign_backup_not_overwritten(tmp_path, job, quote, catalog):
    path = tmp_path / "f2.xlsx"
    openpyxl.Workbook().save(path)
    (tmp_path / "f2.bak.xlsx").write_bytes(b"older backup")
    export_quote(path, job, quote, catalog, append=True)
    assert (tmp_path / "f2.bak.xlsx").read_bytes() == b"older backup"


def test_pcc_workbook_not_backed_up(tmp_path, job, quote, catalog):
    path = tmp_path / "p.xlsx"
    export_quote(path, job, quote, catalog, append=False)
    export_quote(path, job, quote, catalog, append=True)
    assert sorted(p.name for p in tmp_path.iterdir()) == ["p.xlsx"]


@pytest.mark.skipif(__import__("sys").platform == "win32", reason="POSIX modes")
def test_atomic_save_keeps_file_mode(tmp_path, job, quote, catalog):
    import os
    import stat
    path = tmp_path / "m.xlsx"
    export_quote(path, job, quote, catalog, append=False)
    assert stat.S_IMODE(path.stat().st_mode) == 0o644
    os.chmod(path, 0o664)
    export_quote(path, job, quote, catalog, append=True)
    assert stat.S_IMODE(path.stat().st_mode) == 0o664
