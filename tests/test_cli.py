import json
import sys

import openpyxl

from pcc.cli import main
from tests.fixtures.make3mf import MULTI_PLATE_PROJECT


def _file(tmp_path):
    f = tmp_path / "p.3mf"
    f.write_bytes(MULTI_PLATE_PROJECT)
    return f


def test_quote_json_matches_core(cfg_dir, tmp_path, capsys):
    assert main(["quote", str(_file(tmp_path)), "--json"]) == 0
    out = json.loads(capsys.readouterr().out)
    assert set(out) == {"job_id", "currency", "plates", "total_hours", "total_grams",
                        "total_copies", "price"}
    assert len(out["plates"]) == 2


def test_quote_text_and_internal(cfg_dir, tmp_path, capsys):
    assert main(["quote", str(_file(tmp_path)), "--copies", "2", "--internal"]) == 0
    out = capsys.readouterr().out
    assert "×2" in out and "PRICE  ₱" in out and "Internal fee" in out
    assert main(["quote", str(_file(tmp_path))]) == 0
    assert "Internal fee" not in capsys.readouterr().out


def test_quote_export_append(cfg_dir, tmp_path, capsys):
    f, r = _file(tmp_path), tmp_path / "r.xlsx"
    assert main(["quote", str(f), "--export", str(r)]) == 0
    assert main(["quote", str(f), "--export", str(r), "--append"]) == 0
    assert openpyxl.load_workbook(r)["Jobs"].max_row == 3


def test_export_to_invalid_workbook_exit_3(cfg_dir, tmp_path, capsys):
    r = tmp_path / "r.xlsx"
    r.write_text("not a workbook")
    assert main(["quote", str(_file(tmp_path)), "--export", str(r), "--append"]) == 3
    assert "error: Could not write report:" in capsys.readouterr().err


def test_bad_file_exit_3(cfg_dir, tmp_path, capsys):
    f = tmp_path / "bad.3mf"
    f.write_text("nope")
    assert main(["quote", str(f)]) == 3
    assert "Not a valid .3mf file." in capsys.readouterr().err


def test_missing_file_exit_3(cfg_dir, tmp_path, capsys):
    assert main(["parse", str(tmp_path / "none.3mf")]) == 3
    assert capsys.readouterr().err.startswith("error: ")


def test_unknown_printer_exit_2(cfg_dir, tmp_path, capsys):
    assert main(["quote", str(_file(tmp_path)), "--printer", "zzz"]) == 2
    assert "error: " in capsys.readouterr().err


def test_bad_args_exit_2(capsys):
    assert main(["quote"]) == 2
    assert main(["bogus"]) == 2


def test_parse_json(cfg_dir, tmp_path, capsys):
    assert main(["parse", str(_file(tmp_path)), "--json"]) == 0
    out = json.loads(capsys.readouterr().out)
    assert set(out) == {"source_name", "printer_model", "printer_settings_id", "plates"}
    assert set(out["plates"][0]) == {"index", "name", "hours", "filaments"}
    assert main(["parse", str(_file(tmp_path))]) == 0
    assert len(capsys.readouterr().out.strip().splitlines()) == 2


def test_printers_and_filaments(cfg_dir, capsys):
    assert main(["printers", "--json"]) == 0
    assert "Bambu Lab P1S" in [p["name"] for p in json.loads(capsys.readouterr().out)]
    assert main(["filaments", "--json"]) == 0
    assert json.loads(capsys.readouterr().out)
    assert main(["printers"]) == 0 and "P1S" in capsys.readouterr().out
    assert main(["filaments"]) == 0 and "/kg" in capsys.readouterr().out


def test_version(capsys):
    assert main(["--version"]) == 0
    assert capsys.readouterr().out.strip() == "pcc 0.2.0"


def test_serve_without_api(monkeypatch, capsys):
    monkeypatch.setitem(sys.modules, "pcc.api.server", None)
    assert main(["serve"]) == 2
    assert ("API support not installed: pip install print-cost-calculator[api]"
            in capsys.readouterr().err)


def test_append_requires_export(cfg_dir, tmp_path, capsys):
    assert main(["quote", str(_file(tmp_path)), "--append"]) == 2
    assert "--append requires --export" in capsys.readouterr().err


def test_export_refuses_existing_file(cfg_dir, tmp_path, capsys):
    f, r = _file(tmp_path), tmp_path / "r.xlsx"
    r.write_bytes(b"precious")
    for target in (r, tmp_path / "r"):  # .xlsx is added to the second one
        assert main(["quote", str(f), "--export", str(target)]) == 2
        assert ("error: File exists; use --append or choose another path"
                in capsys.readouterr().err)
        assert r.read_bytes() == b"precious"
    assert main(["quote", str(f), "--export", str(r), "--force"]) == 0
    assert openpyxl.load_workbook(r)["Jobs"].max_row == 2


def test_readme_documents_force():
    from pathlib import Path
    readme = (Path(__file__).resolve().parent.parent / "README.md").read_text(encoding="utf-8")
    assert "--force" in readme
