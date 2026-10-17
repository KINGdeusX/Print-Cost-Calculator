"""Appendable Jobs/Plates Excel report (spec section 10)."""

from __future__ import annotations

from pathlib import Path

from pcc.core.currency import excel_money_format
from pcc.core.models import Catalog, Job
from pcc.core.pricing import Quote

JOBS_HEADERS = [
    "Job ID", "Date", "Customer", "Printer", "Currency", "Plates", "Copies",
    "Hours", "Grams", "Filament", "Electricity", "Failure allowance", "Labor",
    "Packaging", "Total cost", "Markup", "Price",
]
PLATES_HEADERS = [
    "Job ID", "Plate #", "Plate", "Source file", "Copies", "Hours/copy",
    "Grams/copy", "Filaments", "Filament", "Electricity", "Failure allowance",
    "Labor", "Plate cost", "Plate price",
]
_DATE_FORMAT = "yyyy-mm-dd hh:mm"
_JOBS_MONEY = range(10, 18)  # columns J..Q
_PLATES_MONEY = range(9, 15)  # columns I..N
_LEGACY = "Legacy v0.1"


def _is_legacy(ws) -> bool:
    vals = {str(c.value).strip().casefold() for c in ws[1] if c.value is not None}
    return {"date", "name", "duration"} <= vals


def _prepare(wb) -> None:
    names = wb.sheetnames
    if "Jobs" not in names and "Plates" not in names and _LEGACY not in names:
        for ws in wb.worksheets:
            if _is_legacy(ws):
                ws.title = _LEGACY
                wb.move_sheet(ws, offset=-wb.index(ws))
                break
    for name, headers in (("Jobs", JOBS_HEADERS), ("Plates", PLATES_HEADERS)):
        if name not in wb.sheetnames:
            from openpyxl.styles import Font

            ws = wb.create_sheet(name)
            ws.append(headers)
            for c in ws[1]:
                c.font = Font(bold=True)
            ws.freeze_panes = "A2"
    for name in ("Sheet", "Sheet1"):
        if name in wb.sheetnames and len(wb.sheetnames) > 2:
            ws = wb[name]
            if ws.max_row == 1 and ws.max_column == 1 and ws["A1"].value is None:
                del wb[name]


def _fit_widths(ws) -> None:
    from openpyxl.utils import get_column_letter

    for col in ws.columns:
        longest = max(
            (len(str(c.value)) if c.value is not None else 0) for c in col
        )
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(longest + 2, 50)


def export_quote(path: Path, job: Job, quote: Quote, catalog: Catalog, append: bool) -> Path:
    import openpyxl

    path = Path(path)
    if path.suffix.lower() != ".xlsx":
        path = path.with_name(path.name + ".xlsx")
    wb = openpyxl.load_workbook(path) if append and path.exists() else openpyxl.Workbook()
    if wb.sheetnames == ["Sheet"]:
        del wb["Sheet"]
    _prepare(wb)

    fmt = excel_money_format(quote.currency_symbol)
    printer = catalog.printer(job.printer_id)
    jobs, plates = wb["Jobs"], wb["Plates"]

    jobs.append([
        job.id, job.created_at, job.customer, printer.name if printer else job.printer_id,
        quote.currency_code, len(quote.plates), quote.total_copies,
        round(quote.total_hours, 2), round(quote.total_grams, 1),
        quote.lines.filament, quote.lines.electricity, quote.lines.failure,
        quote.lines.labor, quote.packaging, quote.job_cost, quote.markup, quote.price,
    ])
    r = jobs.max_row
    jobs.cell(r, 2).number_format = _DATE_FORMAT
    for c in _JOBS_MONEY:
        jobs.cell(r, c).number_format = fmt

    for i, (plate, b) in enumerate(zip(job.plates, quote.plates), start=1):
        parts = []
        for u in plate.filaments:
            fil = catalog.filament(u.filament_id)
            parts.append(f"{fil.name if fil else u.filament_id} {u.grams:.1f} g")
        s = b.per_copy.scaled(b.copies)
        plates.append([
            job.id, i, b.name, plate.source_file or "", b.copies,
            round(b.hours, 2), round(b.grams, 1), "; ".join(parts),
            s.filament, s.electricity, s.failure, s.labor, b.cost_total, b.price_total,
        ])
        for c in _PLATES_MONEY:
            plates.cell(plates.max_row, c).number_format = fmt

    for ws in (jobs, plates):
        _fit_widths(ws)
    wb.save(path)
    return path
